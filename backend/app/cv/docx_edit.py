"""Layout-preserving text editing for .docx documents (AD-5).

Edits rewrite run text (`w:t`) only. Paragraph properties, run styles,
fonts, and document structure are never touched, so the CV's layout is
guaranteed to survive tailoring. Handles target text split across multiple
runs (common when formatting changes mid-word).
"""

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def replace_in_paragraph(
    paragraph: Paragraph,
    old: str,
    new: str,
    replace_all: bool = False,
    ignore_case: bool = False,
) -> bool:
    """Replace `old` with `new` in a paragraph's runs. Returns True on a hit.

    Searching resumes after the inserted text so a replacement containing the
    search string cannot loop forever.
    """
    if not old:
        return False
    replaced = False
    offset = 0
    while True:
        runs = paragraph.runs
        full = "".join(r.text for r in runs)
        haystack = full.lower() if ignore_case else full
        needle = old.lower() if ignore_case else old
        idx = haystack.find(needle, offset)
        if idx == -1:
            break
        _replace_span(runs, idx, idx + len(needle), new)
        replaced = True
        offset = idx + len(new)
        if not replace_all:
            break
    return replaced


def _replace_span(runs: list, start: int, end: int, new_text: str) -> None:
    """Rewrite the text span [start, end) across runs.

    The replacement text lands in the run where the match starts, inheriting
    that run's formatting; fully-covered runs are emptied.
    """
    pos = 0
    start_run = end_run = None
    start_off = end_off = 0
    for run in runs:
        run_len = len(run.text)
        if start_run is None and pos + run_len > start:
            start_run = run
            start_off = start - pos
        if pos + run_len >= end:
            end_run = run
            end_off = end - pos
            break
        pos += run_len

    if start_run is None or end_run is None:
        return

    if start_run is end_run:
        start_run.text = start_run.text[:start_off] + new_text + start_run.text[end_off:]
        return

    start_run.text = start_run.text[:start_off] + new_text
    pos = 0
    clearing = False
    for run in runs:
        run_len = len(run.text)
        if run is start_run:
            clearing = True
            pos += run_len
            continue
        if clearing:
            if run is end_run:
                run.text = run.text[end_off:]
                break
            run.text = ""


def _iter_paragraphs(document: Document):
    yield from document.paragraphs
    yield from _table_paragraphs(document.tables)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            yield from _table_paragraphs(part.tables)


def _table_paragraphs(tables: list[Table]):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                yield from _table_paragraphs(cell.tables)


def replace_in_document(
    document: Document,
    old: str,
    new: str,
    replace_all: bool = True,
    ignore_case: bool = False,
) -> int:
    """Replace across body, tables, headers and footers.

    Returns the number of paragraphs modified.
    """
    count = 0
    for paragraph in _iter_paragraphs(document):
        if replace_in_paragraph(paragraph, old, new, replace_all=replace_all, ignore_case=ignore_case):
            count += 1
    return count
