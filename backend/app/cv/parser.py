"""Read-side parsing of .docx CVs."""

from docx import Document


def extract_text(document: Document) -> str:
    """Extract all visible text (body paragraphs and table cells)."""
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs if p.text.strip())
    return "\n".join(parts)


def extract_sections(document: Document) -> dict[str, str]:
    """Best-effort section split: headings (or bold standalone lines) start a
    section; following paragraphs belong to it until the next heading."""
    sections: dict[str, str] = {}
    current = "_header"
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        is_heading = para.style.name.startswith("Heading") or (
            len(text) < 60 and all(r.bold for r in para.runs if r.text.strip()) and para.runs
        )
        if is_heading:
            current = text
            sections.setdefault(current, "")
        else:
            sections[current] = f"{sections.get(current, '')}\n{text}".strip()
    return sections
