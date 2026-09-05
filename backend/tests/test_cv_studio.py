"""Tests for persistent CV storage, manual content model, and rendering engine."""

import io

from docx import Document
from docx.shared import Pt

from app.cv.content import CVContentData, ExperienceItem, is_empty, to_text
from app.cv.renderer import BODY_SIZE, FONT, NAME_SIZE, render_cv_docx
from app.cv.storage import CVFileStore


class TestCVFileStore:
    def test_save_load_roundtrip(self, tmp_path):
        store = CVFileStore(tmp_path)
        store.save(b"docx-bytes", "my_cv.docx")
        assert store.exists()
        assert store.load() == b"docx-bytes"
        meta = store.meta()
        assert meta["filename"] == "my_cv.docx"
        assert meta["uploaded_at"]
        assert meta["size"] == len(b"docx-bytes")

    def test_load_missing_returns_none(self, tmp_path):
        assert CVFileStore(tmp_path).load() is None
        assert CVFileStore(tmp_path).meta() == {}

    def test_delete(self, tmp_path):
        store = CVFileStore(tmp_path)
        store.save(b"x", "cv.docx")
        store.delete()
        assert not store.exists()
        assert store.meta() == {}

    def test_save_overwrites_previous(self, tmp_path):
        store = CVFileStore(tmp_path)
        store.save(b"old", "old.docx")
        store.save(b"new", "new.docx")
        assert store.load() == b"new"
        assert store.meta()["filename"] == "new.docx"


class TestCVContent:
    def _content(self):
        return CVContentData(
            skills=["Python", "FastAPI"],
            experience=[
                ExperienceItem(
                    role="Backend Engineer",
                    company="Acme",
                    start="Jan 2022",
                    end="Present",
                    description="Built REST services.",
                )
            ],
            achievements=["Won a hackathon"],
        )

    def test_is_empty(self):
        assert is_empty(CVContentData()) is True
        assert is_empty(self._content()) is False
        assert is_empty(CVContentData(skills=["Python"])) is False

    def test_to_text_projects_all_categories(self):
        text = to_text(self._content(), profile_name="Ayesha", education="BSc CS")
        assert "Ayesha" in text
        assert "Skills: Python, FastAPI" in text
        assert "Backend Engineer at Acme (Jan 2022 - Present)" in text
        assert "Built REST services." in text
        assert "- Won a hackathon" in text
        assert "Education: BSc CS" in text


class TestRenderer:
    def _content(self):
        return CVContentData(
            skills=["Python", "FastAPI"],
            experience=[
                ExperienceItem(
                    role="Backend Engineer",
                    company="Acme",
                    start="Jan 2022",
                    end="Present",
                    description="Built REST services.",
                )
            ],
            achievements=["Won a hackathon"],
        )

    def test_renders_all_sections(self):
        data = render_cv_docx(self._content(), name="Ayesha", email="a@b.c", education="BSc CS")
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Ayesha" in text
        assert "a@b.c" in text
        assert "EXPERIENCE" in text
        assert "Backend Engineer · Acme" in text
        assert "Jan 2022 – Present" in text
        assert "SKILLS" in text
        assert "Python, FastAPI" in text
        assert "ACHIEVEMENTS" in text
        assert "Won a hackathon" in text
        assert "EDUCATION" in text

    def test_font_sizes_are_consistent(self):
        """Every run uses the style system's font and one of the fixed sizes."""
        allowed = {NAME_SIZE, Pt(11), Pt(10.5), Pt(10)}
        data = render_cv_docx(self._content(), name="Ayesha", email="a@b.c")
        doc = Document(io.BytesIO(data))
        runs = [run for p in doc.paragraphs for run in p.runs]
        assert runs, "renderer produced no runs"
        for run in runs:
            assert run.font.name == FONT
            assert run.font.size in allowed, f"unexpected size {run.font.size} on '{run.text}'"

    def test_body_text_uses_body_size(self):
        data = render_cv_docx(self._content())
        doc = Document(io.BytesIO(data))
        desc = next(p for p in doc.paragraphs if p.text == "Built REST services.")
        assert desc.runs[0].font.size == BODY_SIZE

    def test_no_empty_paragraphs(self):
        """No blank paragraphs creating excessive whitespace between sections."""
        data = render_cv_docx(self._content(), name="Ayesha")
        doc = Document(io.BytesIO(data))
        assert doc.paragraphs
        for p in doc.paragraphs:
            assert p.text.strip(), "renderer emitted an empty paragraph"

    def test_section_headings_have_fixed_spacing(self):
        data = render_cv_docx(self._content())
        doc = Document(io.BytesIO(data))
        heading = next(p for p in doc.paragraphs if p.text == "EXPERIENCE")
        assert heading.paragraph_format.space_before == Pt(10)
        assert heading.paragraph_format.space_after == Pt(2)

    def test_empty_sections_omitted(self):
        data = render_cv_docx(CVContentData(skills=["Python"]))
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "SKILLS" in text
        assert "EXPERIENCE" not in text
        assert "ACHIEVEMENTS" not in text
        assert "EDUCATION" not in text
