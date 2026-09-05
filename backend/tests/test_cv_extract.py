"""Tests for CV document text extraction and its role in the evidence base."""

import io

from docx import Document

from app.cv.extract import CVTextCache, CVFileStore, extract_docx_text


def _docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, value in enumerate(row):
                table.cell(i, j).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class TestExtractDocxText:
    def test_extracts_paragraphs_and_tables(self):
        data = _docx_bytes(
            ["Ayesha Noman", "Applied AI Engineer"],
            table_rows=[["Python", "3 years"], ["FastAPI", "2 years"]],
        )
        text = extract_docx_text(data)
        assert "Ayesha Noman" in text
        assert "Applied AI Engineer" in text
        assert "Python" in text
        assert "3 years" in text

    def test_skips_blank_and_duplicate_lines(self):
        data = _docx_bytes(["Header", "", "Header", "Body"])
        lines = extract_docx_text(data).splitlines()
        assert lines.count("Header") == 1
        assert "" not in lines

    def test_truncation_enforced_by_cache(self, tmp_path):
        data = _docx_bytes(["word " * 5000])
        cache = CVTextCache(directory=tmp_path, store=CVFileStore(tmp_path))
        cache._store.save(data, "big.docx")
        assert len(cache.get_text()) <= 12000


class TestCVTextCache:
    def test_empty_when_no_file(self, tmp_path):
        cache = CVTextCache(directory=tmp_path, store=CVFileStore(tmp_path))
        assert cache.get_text() == ""

    def test_extracts_and_caches(self, tmp_path):
        store = CVFileStore(tmp_path)
        store.save(_docx_bytes(["Original resume"]), "cv.docx")
        cache = CVTextCache(directory=tmp_path, store=store)
        assert "Original resume" in cache.get_text()
        assert (tmp_path / "cv_text.txt").exists()
        assert (tmp_path / "cv_text_key.txt").exists()

    def test_serves_cache_until_key_changes(self, tmp_path):
        store = CVFileStore(tmp_path)
        store.save(_docx_bytes(["Version one"]), "cv.docx")
        cache = CVTextCache(directory=tmp_path, store=store)
        assert "Version one" in cache.get_text()

        # Overwrite the stored file with a new upload timestamp → new key → re-extract.
        store.save(_docx_bytes(["Version two"]), "cv.docx")
        assert "Version two" in cache.get_text()

    def test_invalidate_removes_cache(self, tmp_path):
        store = CVFileStore(tmp_path)
        store.save(_docx_bytes(["Resume"]), "cv.docx")
        cache = CVTextCache(directory=tmp_path, store=store)
        cache.get_text()
        cache.invalidate()
        assert not (tmp_path / "cv_text.txt").exists()
        assert not (tmp_path / "cv_text_key.txt").exists()
        # Still works after invalidation.
        assert "Resume" in cache.get_text()


class TestCvContextIncludesUploadedDocument:
    def test_uploaded_document_joins_evidence(self, isolated_engine, monkeypatch):
        import app.deps as deps

        monkeypatch.setattr(deps, "get_cv_text", lambda: "Built a RAG pipeline with LangChain.")
        context = deps.cv_context_text()
        assert "UPLOADED CV DOCUMENT" in context
        assert "Built a RAG pipeline with LangChain." in context

    def test_no_file_leaves_document_out(self, isolated_engine, monkeypatch):
        import app.deps as deps

        monkeypatch.setattr(deps, "get_cv_text", lambda: "")
        context = deps.cv_context_text()
        assert "UPLOADED CV DOCUMENT" not in context
