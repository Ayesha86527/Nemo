"""Service-level tests for the CV tailor with a stubbed LLM."""

import io
import json

import pytest
from docx import Document

from app.cv.tailor import CVTailorService, TailorResult, TailoringError
from app.llm.providers import LLMResponse


def _cv_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("I built a REST API with FastAPI and wrote tests.")
    doc.add_paragraph("Skills: Python, SQL")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _stub_generate(plan: dict, error: str | None = None, hint: str | None = None, provider="stub"):
    async def generate(prompt: str, system: str | None = None) -> LLMResponse:
        assert system and "STRICT JSON" in system
        if error:
            return LLMResponse(text="", provider="stub_error", model="m", error=error, hint=hint or "")
        return LLMResponse(text=json.dumps(plan), provider=provider, model="m")

    return generate


class TestCVTailorService:
    async def test_tailor_applies_edits_and_returns_valid_docx(self):
        plan = {
            "summary": "Emphasized API work",
            "edits": [
                {"find": "REST API with FastAPI", "replace": "high-throughput REST API with FastAPI"},
                {"find": "Skills: Python, SQL", "replace": "Skills: Python, SQL, Docker"},
            ],
        }
        service = CVTailorService(generate_fn=_stub_generate(plan))
        result = await service.tailor(_cv_bytes(), "Backend engineer, FastAPI, Docker")

        assert isinstance(result, TailorResult)
        assert result.edits_applied == 2
        assert result.edits_skipped == 0
        reopened = Document(io.BytesIO(result.output))
        text = "\n".join(p.text for p in reopened.paragraphs)
        assert "high-throughput REST API" in text
        assert "Docker" in text

    async def test_tailor_wraps_markdown_fenced_json(self):
        plan_json = json.dumps({"summary": "s", "edits": [{"find": "Jane Doe", "replace": "Jane A. Doe"}]})

        async def generate(prompt, system=None):
            return LLMResponse(text=f"```json\n{plan_json}\n```", provider="stub", model="m")

        service = CVTailorService(generate_fn=generate)
        result = await service.tailor(_cv_bytes(), "any job")
        assert result.edits_applied == 1

    async def test_invalid_docx_raises_with_hint(self):
        service = CVTailorService(generate_fn=_stub_generate({}))
        with pytest.raises(TailoringError) as exc_info:
            await service.tailor(b"not a docx", "job")
        assert exc_info.value.hint

    async def test_llm_failure_propagates_as_tailoring_error(self):
        service = CVTailorService(
            generate_fn=_stub_generate({}, error="Ollama is not reachable.", hint="Start Ollama.")
        )
        with pytest.raises(TailoringError) as exc_info:
            await service.tailor(_cv_bytes(), "job")
        assert "Ollama" in str(exc_info.value)
        assert exc_info.value.hint == "Start Ollama."

    async def test_unmatched_edits_counted_as_skipped(self):
        plan = {
            "summary": "",
            "edits": [
                {"find": "text that does not exist", "replace": "x"},
                {"find": "Jane Doe", "replace": "Jane Doe-Strong"},
            ],
        }
        service = CVTailorService(generate_fn=_stub_generate(plan))
        result = await service.tailor(_cv_bytes(), "job")
        assert result.edits_applied == 1
        assert result.edits_skipped == 1

    async def test_all_edits_unmatched_raises(self):
        plan = {"summary": "", "edits": [{"find": "nothing matches", "replace": "x"}]}
        service = CVTailorService(generate_fn=_stub_generate(plan))
        with pytest.raises(TailoringError):
            await service.tailor(_cv_bytes(), "job")

    async def test_case_insensitive_fallback(self):
        plan = {"summary": "", "edits": [{"find": "JANE DOE", "replace": "Jane Doe, Engineer"}]}
        service = CVTailorService(generate_fn=_stub_generate(plan))
        result = await service.tailor(_cv_bytes(), "job")
        assert result.edits_applied == 1

    async def test_research_context_injected_when_provided(self):
        seen_prompts = []

        async def generate(prompt, system=None):
            seen_prompts.append(prompt)
            return LLMResponse(
                text=json.dumps({"summary": "", "edits": [{"find": "Jane Doe", "replace": "Jane X"}]}),
                provider="stub",
                model="m",
            )

        service = CVTailorService(generate_fn=generate)
        await service.tailor(_cv_bytes(), "ML job", research="Company builds ML pipelines at scale.")
        assert "Company builds ML pipelines at scale." in seen_prompts[0]
