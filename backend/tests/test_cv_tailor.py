"""Tests for layout-preserving .docx text editing (AD-5).

Core invariant: replacements rewrite run text only — styles, fonts, and
paragraph properties are never touched. Text split across multiple runs
(e.g. a bold word in the middle of a sentence) must still match.
"""

import io

from docx import Document
from docx.shared import Pt

from app.cv.docx_edit import replace_in_document, replace_in_paragraph


def _doc_with_paragraph(text: str) -> Document:
    doc = Document()
    doc.add_paragraph(text)
    return doc


def _doc_with_split_runs() -> Document:
    """Paragraph whose text is split across runs with mixed formatting:
    'Built a REST API with FastAPI' as ['Built a ', 'REST API'(bold), ' with FastAPI']"""
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Built a ")
    bold = para.add_run("REST API")
    bold.bold = True
    para.add_run(" with FastAPI")
    return doc


class TestReplaceInParagraph:
    def test_single_run_replacement(self):
        doc = _doc_with_paragraph("I know Python well.")
        para = doc.paragraphs[0]
        assert replace_in_paragraph(para, "Python", "Go") is True
        assert para.text == "I know Go well."

    def test_replacement_spanning_multiple_runs(self):
        doc = _doc_with_split_runs()
        para = doc.paragraphs[0]
        assert replace_in_paragraph(para, "REST API with FastAPI", "GraphQL API") is True
        assert para.text == "Built a GraphQL API"

    def test_no_match_returns_false(self):
        doc = _doc_with_paragraph("Nothing to see here.")
        assert replace_in_paragraph(doc.paragraphs[0], "absent", "x") is False

    def test_formatting_preserved_on_first_run(self):
        doc = _doc_with_split_runs()
        para = doc.paragraphs[0]
        replace_in_paragraph(para, "REST API", "GraphQL")
        runs = [r for r in para.runs if r.text]
        assert any(r.bold for r in runs), "bold formatting must survive the edit"

    def test_paragraph_style_preserved(self):
        doc = Document()
        para = doc.add_paragraph("Heading text", style="Heading 1")
        replace_in_paragraph(para, "Heading text", "New heading")
        assert para.style.name == "Heading 1"
        assert para.text == "New heading"

    def test_replacement_at_paragraph_start_and_end(self):
        doc = _doc_with_paragraph("old start middle end old")
        para = doc.paragraphs[0]
        replace_in_paragraph(para, "old start", "new start")
        assert para.text.startswith("new start")
        replace_in_paragraph(para, "end old", "end new")
        assert para.text.endswith("end new")

    def test_multiple_occurrences_in_one_paragraph(self):
        doc = _doc_with_paragraph("Python and Python again")
        para = doc.paragraphs[0]
        replace_in_paragraph(para, "Python", "Go", replace_all=True)
        assert para.text == "Go and Go again"

    def test_case_insensitive_match(self):
        doc = _doc_with_paragraph("Expert in PYTHON.")
        para = doc.paragraphs[0]
        assert replace_in_paragraph(para, "python", "Go", ignore_case=True) is True
        assert para.text == "Expert in Go."

    def test_replacement_containing_search_string_terminates(self):
        doc = _doc_with_paragraph("REST API work")
        para = doc.paragraphs[0]
        replace_in_paragraph(para, "REST API", "high-throughput REST API", replace_all=True)
        assert para.text == "high-throughput REST API work"


class TestReplaceInDocument:
    def test_replaces_across_paragraphs(self):
        doc = Document()
        doc.add_paragraph("I love Java.")
        doc.add_paragraph("Java is my life.")
        count = replace_in_document(doc, "Java", "Rust", replace_all=True)
        assert count == 2
        assert doc.paragraphs[0].text == "I love Rust."
        assert doc.paragraphs[1].text == "Rust is my life."

    def test_replaces_inside_tables(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Skill"
        table.cell(0, 1).text = "Java"
        count = replace_in_document(doc, "Java", "Kotlin")
        assert count == 1
        assert table.cell(0, 1).text == "Kotlin"

    def test_zero_when_absent(self):
        doc = _doc_with_paragraph("hello")
        assert replace_in_document(doc, "zzz", "yyy") == 0

    def test_roundtrip_stays_valid_docx(self):
        doc = _doc_with_split_runs()
        replace_in_document(doc, "REST API", "GraphQL")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        reopened = Document(buf)
        assert reopened.paragraphs[0].text == "Built a GraphQL with FastAPI"


class TestParser:
    def test_extract_text_includes_tables(self):
        from app.cv.parser import extract_text

        doc = Document()
        doc.add_paragraph("Summary paragraph.")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Cell content"
        text = extract_text(doc)
        assert "Summary paragraph." in text
        assert "Cell content" in text
